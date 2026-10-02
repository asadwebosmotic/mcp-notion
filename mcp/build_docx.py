"""Convert KNOWLEDGE_TRANSFER.md into a beautifully formatted Word (.docx) document."""

import re
import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Set the background color of a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tc_pr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set cell padding (in dxa: 20 dxa = 1 pt)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)

def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    """Set custom borders on a cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
    borders = {'top': top, 'bottom': bottom, 'left': left, 'right': right}
    for side, border in borders.items():
        if border:
            el = parse_xml(f'<w:{side} {nsdecls("w")} w:val="{border.get("val", "single")}" w:sz="{border.get("sz", 4)}" w:space="0" w:color="{border.get("color", "CCCCCC")}"/>')
            tc_borders.append(el)
        else:
            el = parse_xml(f'<w:{side} {nsdecls("w")} w:val="none"/>')
            tc_borders.append(el)
    tc_pr.append(tc_borders)

def add_code_block(doc, code_text):
    """Add a shaded code block with a clean monospace font and border."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    set_cell_borders(cell, 
                     left={'val': 'single', 'sz': 16, 'color': '0284C7'},
                     top={'val': 'single', 'sz': 4, 'color': 'E2E8F0'},
                     bottom={'val': 'single', 'sz': 4, 'color': 'E2E8F0'},
                     right={'val': 'single', 'sz': 4, 'color': 'E2E8F0'})
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(code_text.strip())
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(30, 41, 59)
    
    # Empty spacing after table
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(4)

def add_callout_box(doc, text, title="NOTE"):
    """Add an alert/callout box with colored left bar."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "EFF6FF")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    set_cell_borders(cell, 
                     left={'val': 'single', 'sz': 24, 'color': '2563EB'},
                     top={'val': 'single', 'sz': 4, 'color': 'DBEAFE'},
                     bottom={'val': 'single', 'sz': 4, 'color': 'DBEAFE'},
                     right={'val': 'single', 'sz': 4, 'color': 'DBEAFE'})
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.2
    run_title = p.add_run(f"[{title}] ")
    run_title.font.name = "Segoe UI"
    run_title.font.bold = True
    run_title.font.size = Pt(9.5)
    run_title.font.color.rgb = RGBColor(37, 99, 235)
    
    run_text = p.add_run(text)
    run_text.font.name = "Segoe UI"
    run_text.font.size = Pt(9.5)
    run_text.font.color.rgb = RGBColor(30, 58, 138)
    
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(4)

def format_inline_runs(paragraph, text, base_font="Segoe UI", base_size=Pt(10), base_color=RGBColor(30, 41, 59)):
    """Parse bold, inline code, and normal text into styled runs."""
    pattern = re.compile(r'(\*\*.*?\*\*|`.*?`|\[.*?\]\(.*?\))')
    parts = pattern.split(text)
    
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.font.name = base_font
            run.font.size = base_size
            run.font.bold = True
            run.font.color.rgb = base_color
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(f" {part[1:-1]} ")
            run.font.name = "Consolas"
            run.font.size = Pt(base_size.pt * 0.9)
            run.font.color.rgb = RGBColor(185, 28, 28)
        elif part.startswith("[") and "]" in part and "(" in part and part.endswith(")"):
            m = re.match(r'\[(.*?)\]\((.*?)\)', part)
            if m:
                label, url = m.groups()
                run = paragraph.add_run(label)
                run.font.name = base_font
                run.font.size = base_size
                run.font.color.rgb = RGBColor(37, 99, 235)
                run.font.underline = True
            else:
                run = paragraph.add_run(part)
                run.font.name = base_font
                run.font.size = base_size
                run.font.color.rgb = base_color
        else:
            run = paragraph.add_run(part)
            run.font.name = base_font
            run.font.size = base_size
            run.font.color.rgb = base_color

def build_docx(md_path, docx_path):
    with open(md_path, "r", encoding="utf-8") as f:
        md_content = f.read()

    doc = Document()
    
    # Page Margins: 0.8 inch
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
    
    # Set default style
    style_normal = doc.styles['Normal']
    font_normal = style_normal.font
    font_normal.name = 'Segoe UI'
    font_normal.size = Pt(10)
    font_normal.color.rgb = RGBColor(30, 41, 59)
    
    lines = md_content.splitlines()
    i = 0
    n = len(lines)
    
    # Header Banner
    header_table = doc.add_table(rows=1, cols=1)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.autofit = False
    h_cell = header_table.cell(0, 0)
    h_cell.width = Inches(6.9)
    set_cell_background(h_cell, "0F172A")
    set_cell_margins(h_cell, top=260, bottom=260, left=260, right=260)
    set_cell_borders(h_cell)
    
    hp = h_cell.paragraphs[0]
    hp.paragraph_format.space_before = Pt(0)
    hp.paragraph_format.space_after = Pt(4)
    hrun1 = hp.add_run("KNOWLEDGE TRANSFER DOCUMENT")
    hrun1.font.name = "Segoe UI"
    hrun1.font.size = Pt(20)
    hrun1.font.bold = True
    hrun1.font.color.rgb = RGBColor(248, 250, 252)
    
    hp2 = h_cell.add_paragraph()
    hp2.paragraph_format.space_before = Pt(0)
    hp2.paragraph_format.space_after = Pt(8)
    hrun2 = hp2.add_run("Notion, Slack & Google Workspace MCP Assistant & Orchestrator")
    hrun2.font.name = "Segoe UI"
    hrun2.font.size = Pt(13)
    hrun2.font.color.rgb = RGBColor(147, 197, 253)
    
    hp3 = h_cell.add_paragraph()
    hp3.paragraph_format.space_before = Pt(4)
    hp3.paragraph_format.space_after = Pt(0)
    hrun3 = hp3.add_run("Git Repository: https://github.com/asadintwala/mcp  •  Version 1.0.0  •  Target: Engineering & Technical Leads")
    hrun3.font.name = "Segoe UI"
    hrun3.font.size = Pt(9)
    hrun3.font.color.rgb = RGBColor(203, 213, 225)
    
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(12)
    
    in_code_block = False
    code_block_lines = []
    code_block_lang = ""
    
    while i < n:
        line = lines[i]
        
        # Skip the original document title lines since we rendered the executive banner
        if i == 0 and line.startswith("# Knowledge Transfer"):
            i += 1
            while i < n and (lines[i].startswith("**Repository:") or lines[i].startswith("**Document") or lines[i].startswith("**Target") or lines[i].startswith("**Last") or lines[i] == "---" or lines[i].strip() == ""):
                i += 1
            continue
            
        # Code block boundary
        if line.startswith("```"):
            if not in_code_block:
                in_code_block = True
                code_block_lang = line[3:].strip()
                code_block_lines = []
            else:
                in_code_block = False
                code_text = "\n".join(code_block_lines)
                add_code_block(doc, code_text)
            i += 1
            continue
            
        if in_code_block:
            code_block_lines.append(line)
            i += 1
            continue
            
        # Markdown table detection
        if line.strip().startswith("|") and line.strip().endswith("|"):
            table_lines = []
            while i < n and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            
            # Parse table
            if len(table_lines) >= 2:
                header_row = [c.strip() for c in table_lines[0].split("|")[1:-1]]
                data_rows = []
                for t_line in table_lines[2:]: # skip separator line
                    data_rows.append([c.strip() for c in t_line.split("|")[1:-1]])
                
                table = doc.add_table(rows=len(data_rows) + 1, cols=len(header_row))
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.autofit = False
                
                # Style Header Row
                for col_idx, text in enumerate(header_row):
                    cell = table.cell(0, col_idx)
                    set_cell_background(cell, "1E293B")
                    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                    set_cell_borders(cell, 
                                     top={'val': 'single', 'sz': 4, 'color': '334155'},
                                     bottom={'val': 'single', 'sz': 12, 'color': '0F172A'},
                                     left={'val': 'single', 'sz': 4, 'color': '334155'},
                                     right={'val': 'single', 'sz': 4, 'color': '334155'})
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    run = p.add_run(text)
                    run.font.name = "Segoe UI"
                    run.font.bold = True
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = RGBColor(255, 255, 255)
                
                # Style Data Rows
                for row_idx, row_data in enumerate(data_rows):
                    bg_color = "FFFFFF" if row_idx % 2 == 0 else "F8FAFC"
                    for col_idx in range(len(header_row)):
                        cell = table.cell(row_idx + 1, col_idx)
                        text = row_data[col_idx] if col_idx < len(row_data) else ""
                        set_cell_background(cell, bg_color)
                        set_cell_margins(cell, top=80, bottom=80, left=140, right=140)
                        set_cell_borders(cell, 
                                         top={'val': 'single', 'sz': 4, 'color': 'E2E8F0'},
                                         bottom={'val': 'single', 'sz': 4, 'color': 'E2E8F0'},
                                         left={'val': 'single', 'sz': 4, 'color': 'E2E8F0'},
                                         right={'val': 'single', 'sz': 4, 'color': 'E2E8F0'})
                        p = cell.paragraphs[0]
                        p.paragraph_format.space_before = Pt(1)
                        p.paragraph_format.space_after = Pt(1)
                        format_inline_runs(p, text, base_font="Segoe UI", base_size=Pt(9), base_color=RGBColor(51, 65, 85))
                
                # Spacing after table
                spacer = doc.add_paragraph()
                spacer.paragraph_format.space_before = Pt(0)
                spacer.paragraph_format.space_after = Pt(6)
            continue
            
        # Horizontal Rule
        if line.strip() in ["---", "***", "___"]:
            # Subtle divider
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="CBD5E1"/></w:pBdr>')
            p._p.get_or_add_pPr().append(pBdr)
            i += 1
            continue
            
        # Headings
        if line.startswith("# "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(line[2:].strip())
            run.font.name = "Segoe UI"
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = RGBColor(15, 23, 42)
            i += 1
            continue
            
        if line.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(line[3:].strip())
            run.font.name = "Segoe UI"
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = RGBColor(30, 58, 138)
            i += 1
            continue
            
        if line.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(line[4:].strip())
            run.font.name = "Segoe UI"
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = RGBColor(51, 65, 85)
            i += 1
            continue
            
        # Bullet list item
        if line.strip().startswith("- ") or line.strip().startswith("* "):
            bullet_indent = len(line) - len(line.lstrip())
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            content = line.strip()[2:]
            format_inline_runs(p, content, base_font="Segoe UI", base_size=Pt(9.5), base_color=RGBColor(30, 41, 59))
            i += 1
            continue
            
        # Numbered list item
        numbered_match = re.match(r'^\s*(\d+)\.\s+(.*)$', line)
        if numbered_match:
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            content = numbered_match.group(2)
            format_inline_runs(p, content, base_font="Segoe UI", base_size=Pt(9.5), base_color=RGBColor(30, 41, 59))
            i += 1
            continue
            
        # Standard paragraph
        if line.strip():
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.18
            format_inline_runs(p, line.strip(), base_font="Segoe UI", base_size=Pt(9.5), base_color=RGBColor(30, 41, 59))
        
        i += 1

    doc.save(docx_path)
    print(f"Successfully generated DOCX at: {docx_path}")

if __name__ == "__main__":
    md_file = r"d:\My Pocs\mcp-notion\KNOWLEDGE_TRANSFER.md"
    docx_file = r"d:\My Pocs\mcp-notion\KNOWLEDGE_TRANSFER.docx"
    build_docx(md_file, docx_file)
